from config import get_config
from Initialization import Init
import pandas as pd
import openpyxl as op

# state_list = pd.read_excel("../Excel/state-2.xlsx", nrows=8225, sheet_name='Sheet1', index_col=False, header=None)
state_list = pd.read_excel("Excel/state-2.xlsx", nrows=8225, sheet_name='Sheet1', index_col=False, header=None)
action_space = pd.read_excel('Excel/action space 21.xlsx', index_col=False, header=None)

def ModeltoExcel(config):
    number_index = 0
    wb = op.Workbook()  # 创建工作簿对象
    ws = wb['Sheet']  # 创建子表

    initialization = Init(config)  # 初始化

    agent = initialization.init_agent()  # 初始化智能体

    agent.load()    # loading model

    for i in range(8225):

        action = agent.get_action(i)
        a = action_space.values[0, action]
        ws.cell(row=number_index + 1, column=1).value = a

        print('action:', a)

        # Mean, Std = agent.ModelToMeanAndStd(i)
        #
        # Mean = float(Mean)
        #
        # Std = float(Std)
        #
        # ws.cell(row=number_index + 1, column=1).value = Mean
        #
        # ws.cell(row=number_index + 1, column=2).value = Std

        wb.save('Train_Excel/ActionM_DQN.12.3.xlsx')
        number_index = number_index + 1

        # print('Mean:', Mean)

        # print('Std:', Std)

        print('This is number index:', i)

if __name__ == '__main__':
    parser = get_config()
    config = parser.parse_args()
    ModeltoExcel(config)